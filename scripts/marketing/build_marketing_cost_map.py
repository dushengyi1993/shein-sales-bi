from __future__ import annotations

"""Build the local cost map used by marketing campaign signup helpers.

The output stays under tmp/ because it is generated operating data, while this
script is committed so the workflow can be reproduced on another machine.
"""

import json
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path


ROOT = Path(__file__).resolve().parents[2]
COST_INPUT_CANDIDATES = [
    ROOT / "inputs" / "costs" / "成本.xlsx",
    ROOT / "inputs" / "costs" / "成本计算表.xlsx",
]
OUTPUT = ROOT / "tmp" / "mbrs" / "marketing-cost-map.json"
BI_PATH = ROOT / "outputs" / "bi-portal" / "data.json"
BI_PROFIT_SECTION_PATH = ROOT / "outputs" / "bi-portal" / "sections" / "profit.json"
BI_INVENTORY_SECTION_PATH = ROOT / "outputs" / "bi-portal" / "sections" / "inventoryTrend.json"
CNY_TO_SAR = 1 / 1.8
MAX_STORAGE_INVENTORY_DATE_GAP_DAYS = 2
MAX_STORAGE_INVENTORY_QUANTITY_RELATIVE_DIFFERENCE = 0.15


def clean(value) -> str:
    return "" if value is None else str(value).strip()


def num(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = re.sub(r"[^0-9.\-]", "", str(value).strip().replace(",", ""))
    if s in {"", "-", ".", "-."}:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def compact(value: object) -> str:
    return "".join(ch for ch in clean(value).upper() if ch.isalnum())


def lookup_cost(cost_map: dict[str, float], keys: list[str]) -> float | None:
    by_compact = {compact(k): v for k, v in cost_map.items()}
    for key in keys:
        if key in cost_map:
            return cost_map[key]
        c = compact(key)
        if c in by_compact:
            return by_compact[c]
    return None


def parse_date(value):
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value)[:10]).date()
    except ValueError:
        return None


def parse_timestamp(value):
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return parsed if parsed.tzinfo is not None else parsed.replace(tzinfo=timezone.utc)
    except ValueError:
        return None


def extract_section(document: dict, key: str) -> dict | None:
    if not isinstance(document, dict):
        return None
    data = document.get("data") if isinstance(document.get("data"), dict) else document
    section = data.get(key) if isinstance(data, dict) else None
    return section if isinstance(section, dict) else None


def prefer_section_payload(bi: dict, section_document: dict, key: str) -> bool:
    """Use a targeted section when it is at least as fresh as full data.json."""

    section = extract_section(section_document, key)
    if section is None:
        return False
    if not isinstance(bi.get(key), dict) or not bi[key]:
        return True
    full_generated_at = parse_timestamp(bi.get("generatedAt"))
    section_generated_at = parse_timestamp(section_document.get("generatedAt"))
    if section_generated_at is None:
        return full_generated_at is None
    if full_generated_at is None:
        return True
    return section_generated_at >= full_generated_at


def build_operational_inventory_map(bi: dict) -> dict[str, dict]:
    """Index ET inventory used only to cross-check the billing denominator."""

    out: dict[str, dict] = {}
    rows = (bi.get("inventoryDepletion") or {}).get("products") or []
    for row in rows:
        standard = clean(row.get("standard_goods_sn"))
        key = compact(standard or row.get("match_key"))
        if not key:
            continue
        match_status = clean(row.get("inventory_match_status")).lower()
        sellable = num(
            row.get("current_sellable_quantity")
            if row.get("current_sellable_quantity") is not None
            else row.get("et_estimated_available_qty")
        )
        damaged = num(row.get("et_damaged_qty"))
        damaged = max(float(damaged or 0.0), 0.0)
        physical_quantity = None
        if match_status == "matched" and sellable is not None and sellable >= 0:
            physical_quantity = max(float(sellable), 0.0) + damaged
        operational_stock_policy = clean(row.get("et_operational_stock_policy"))
        store_snapshot_date = parse_date(row.get("et_store_snapshot_date"))
        box_snapshot_date = parse_date(row.get("et_box_snapshot_date"))
        if "full_carton" in operational_stock_policy.lower():
            # A blended loose + full-carton quantity is only as fresh as its
            # older contributing snapshot.  Do not let a fresh box snapshot
            # conceal stale loose-stock evidence (or vice versa).
            required_dates = [store_snapshot_date, box_snapshot_date]
            snapshot_date = min(required_dates) if all(required_dates) else None
        else:
            # Loose-only operational quantity comes from the store snapshot.
            # A fresh box snapshot must not conceal missing/stale loose stock.
            snapshot_date = store_snapshot_date
        evidence = {
            "standard": standard,
            "matchStatus": match_status or "missing",
            "sellableQuantity": None if sellable is None else round(max(float(sellable), 0.0), 4),
            "damagedQuantity": round(damaged, 4),
            "physicalQuantity": None if physical_quantity is None else round(physical_quantity, 4),
            "snapshotDate": snapshot_date,
            "operationalStockPolicy": operational_stock_policy,
        }
        for candidate in {key, compact(standard), compact(row.get("match_key"))}:
            if candidate:
                out[candidate] = evidence
    return out


def storage_inventory_dates_are_compatible(storage_date, inventory_date) -> tuple[bool, int | None]:
    if storage_date is None or inventory_date is None:
        return False, None
    gap_days = abs((storage_date - inventory_date).days)
    return gap_days <= MAX_STORAGE_INVENTORY_DATE_GAP_DAYS, gap_days


def build_storage_unit_map(bi: dict) -> dict[str, dict]:
    """Build storage cost per sellable unit from ET product storage daily rows.

    `profit.products.storage_fee_sar` is an all-history aggregate. It must not
    be divided by historical sold quantity for campaign pricing: slow-moving
    products with one or two sales would get a wildly inflated "per unit"
    storage cost. The ET daily detail already has inventory quantity per day.
    Campaign pricing should use the moving-average accumulated storage cost
    carried by the inventory that is still in storage: each day adds that
    day's storage fee to the inventory cost balance; when the charged quantity
    drops, the sold/outbound units take away their proportional accumulated
    balance. This removes storage cost that belongs to units no longer in
    storage while avoiding unsupported batch/serial assumptions.  The billing
    detail is unitized upstream (including box-item CaseQuantity fallback) and
    remains the denominator because its scope matches the fee numerator.  A
    fresh, matched operational sellable + damaged quantity is only a bounded
    cross-check: stale evidence or a relative mismatch above 15% makes the
    storage unit cost unavailable instead of silently diluting it with newly
    arrived or otherwise unbilled stock.
    """

    rows = (bi.get("profit") or {}).get("productStorageDaily") or []
    operational_inventory_by_key = build_operational_inventory_map(bi)
    dated = [parse_date(r.get("date")) for r in rows]
    max_date = max((d for d in dated if d), default=None)
    cutoff = max_date - timedelta(days=29) if max_date else None
    rows_by_key: dict[str, list[dict]] = {}

    for row in rows:
        standard = clean(row.get("standard_goods_sn"))
        key = compact(standard)
        if not key:
            continue
        rows_by_key.setdefault(key, []).append(row)

    by_key: dict[str, dict] = {}
    for key, sku_rows in rows_by_key.items():
        sku_rows = sorted(sku_rows, key=lambda r: (parse_date(r.get("date")) or datetime.min.date()).isoformat())
        rec = by_key.setdefault(key, {
            "standard": clean(sku_rows[-1].get("standard_goods_sn")) or clean(sku_rows[0].get("standard_goods_sn")),
            "totalFeeSar": 0.0,
            "allUnitSar": 0.0,
            "allDays": 0,
            "recent30FeeSar": 0.0,
            "recent30UnitSar": 0.0,
            "recent30Days": 0,
            "quantityDays": 0.0,
            "currentQuantity": None,
            "inventoryCostBalanceSar": 0.0,
            "nonpositiveQuantityFeeDays": 0,
            "methods": set(),
            "maxDate": None,
            "minDate": None,
        })
        previous_quantity = 0.0
        inventory_cost_balance = 0.0
        for row in sku_rows:
            fee = num(row.get("storage_fee_sar"))
            quantity = num(row.get("storage_quantity"))
            if quantity is None:
                quantity = num(row.get("quantity"))
            unit = num(row.get("storage_fee_per_unit_sar"))
            if unit is None and fee is not None and quantity and quantity > 0:
                unit = fee / quantity
            day = parse_date(row.get("date"))
            method = clean(row.get("storage_fee_method"))
            if method:
                rec["methods"].add(method)

            if quantity is not None:
                current_quantity = max(float(quantity), 0.0)
                if previous_quantity > 0 and current_quantity < previous_quantity and inventory_cost_balance > 0:
                    outbound_quantity = previous_quantity - current_quantity
                    previous_average = inventory_cost_balance / previous_quantity
                    inventory_cost_balance = max(0.0, inventory_cost_balance - previous_average * outbound_quantity)
                previous_quantity = current_quantity
                rec["currentQuantity"] = current_quantity

            if fee is not None and abs(float(fee)) > 1e-9 and (quantity is None or quantity <= 0):
                rec["nonpositiveQuantityFeeDays"] += 1

            if fee is not None:
                rec["totalFeeSar"] += fee
                if previous_quantity > 0:
                    inventory_cost_balance += fee
            if quantity:
                rec["quantityDays"] += quantity
            if unit is not None:
                rec["allUnitSar"] += unit
                rec["allDays"] += 1
            if day:
                rec["minDate"] = day if rec["minDate"] is None else min(rec["minDate"], day)
                rec["maxDate"] = day if rec["maxDate"] is None else max(rec["maxDate"], day)
            if cutoff and day and day >= cutoff:
                if fee is not None:
                    rec["recent30FeeSar"] += fee
                if unit is not None:
                    rec["recent30UnitSar"] += unit
                    rec["recent30Days"] += 1
        rec["inventoryCostBalanceSar"] = inventory_cost_balance

    out: dict[str, dict] = {}
    for key, rec in by_key.items():
        billed_current_quantity = rec["currentQuantity"] or 0.0
        inventory = operational_inventory_by_key.get(key)
        operational_quantity = num(inventory.get("physicalQuantity")) if inventory else None
        inventory_compatible, date_gap_days = storage_inventory_dates_are_compatible(
            rec["maxDate"],
            inventory.get("snapshotDate") if inventory else None,
        )

        candidate_unit = (
            rec["inventoryCostBalanceSar"] / billed_current_quantity
            if billed_current_quantity > 0
            else None
        )
        quantity_ratio = (
            operational_quantity / billed_current_quantity
            if operational_quantity is not None and billed_current_quantity > 0
            else None
        )
        quantity_relative_difference = (
            abs(operational_quantity - billed_current_quantity)
            / max(operational_quantity, billed_current_quantity)
            if operational_quantity is not None
            and operational_quantity > 0
            and billed_current_quantity > 0
            else None
        )

        allocation_quantity = None
        allocation_source = "blocked_missing_inventory_crosscheck"
        evidence_status = "missing_inventory_crosscheck"
        if rec["nonpositiveQuantityFeeDays"] > 0:
            allocation_source = "blocked_historical_nonpositive_quantity_with_fee"
            evidence_status = "historical_nonpositive_quantity_with_fee"
        elif inventory:
            if inventory.get("matchStatus") != "matched" or operational_quantity is None:
                allocation_source = "blocked_inventory_not_fresh_matched"
                evidence_status = "inventory_not_fresh_matched"
            elif not inventory_compatible:
                allocation_source = "blocked_inventory_storage_date_mismatch"
                evidence_status = "inventory_storage_date_mismatch"
            elif billed_current_quantity <= 0 or operational_quantity <= 0:
                allocation_source = "blocked_nonpositive_quantity"
                evidence_status = "nonpositive_quantity_evidence"
            elif quantity_relative_difference is None or quantity_relative_difference > MAX_STORAGE_INVENTORY_QUANTITY_RELATIVE_DIFFERENCE:
                allocation_source = "blocked_inventory_storage_quantity_mismatch"
                evidence_status = "inventory_storage_quantity_mismatch"
            else:
                allocation_quantity = billed_current_quantity
                allocation_source = "unitized_storage_daily_billed_current_quantity"
                evidence_status = "fresh_quantity_crosscheck_passed"

        unit = candidate_unit if allocation_quantity is not None else None
        storage_unit_basis = (
            "moving_average_remaining_inventory_storage_cost_over_unitized_billed_inventory"
            if unit is not None
            else "suspect_quantity_evidence_fail_closed"
        )
        out[key] = {
            "standard": rec["standard"],
            "storageUnitCostSar": None if unit is None else round(unit, 4),
            "storageUnitCostCandidateSar": None if candidate_unit is None else round(candidate_unit, 4),
            "storageRecent30UnitCostSar": round(rec["recent30UnitSar"], 4) if rec["recent30Days"] else None,
            "storageAllHistoryUnitCostSar": round(rec["allUnitSar"], 4) if rec["allDays"] else None,
            "storageFeeSar": round(rec["totalFeeSar"], 4),
            "storageInventoryCostBalanceSar": round(rec["inventoryCostBalanceSar"], 4),
            # Backward-compatible field: this remains the current quantity in
            # the storage-fee detail, not the widened shared allocation pool.
            "storageCurrentQuantity": round(billed_current_quantity, 4),
            "storageBilledCurrentQuantity": round(billed_current_quantity, 4),
            "storageOperationalSellableQuantity": None if not inventory else inventory.get("sellableQuantity"),
            "storageOperationalDamagedQuantity": None if not inventory else inventory.get("damagedQuantity"),
            "storageOperationalPhysicalQuantity": None if operational_quantity is None else round(operational_quantity, 4),
            "storageOperationalInventorySnapshotDate": None if not inventory or inventory.get("snapshotDate") is None else inventory["snapshotDate"].isoformat(),
            "storageOperationalStockPolicy": "" if not inventory else inventory.get("operationalStockPolicy"),
            "storageAllocationQuantity": None if allocation_quantity is None else round(allocation_quantity, 4),
            "storageAllocationQuantitySource": allocation_source,
            "storageQuantityEvidenceStatus": evidence_status,
            "storageQuantityDateGapDays": date_gap_days,
            "storageQuantityRatioOperationalToBilled": None if quantity_ratio is None else round(quantity_ratio, 4),
            "storageQuantityRelativeDifference": None if quantity_relative_difference is None else round(quantity_relative_difference, 4),
            "storageNonpositiveQuantityFeeDays": rec["nonpositiveQuantityFeeDays"],
            "storageRecent30FeeSar": round(rec["recent30FeeSar"], 4),
            "storageQuantityDays": round(rec["quantityDays"], 4),
            "storageRecent30Days": rec["recent30Days"],
            "storageAllDays": rec["allDays"],
            "storageMethod": " / ".join(sorted(rec["methods"])) or "missing",
            "storageSourceDateMin": rec["minDate"].isoformat() if rec["minDate"] else None,
            "storageSourceDateMax": rec["maxDate"].isoformat() if rec["maxDate"] else None,
            "storageUnitBasis": storage_unit_basis,
        }
    return out


def build_true_cost_map(cost_map: dict[str, float]) -> dict[str, dict]:
    if not BI_PATH.exists():
        return {}
    try:
        bi = json.loads(BI_PATH.read_text(encoding="utf-8"))
    except Exception:
        return {}
    if BI_PROFIT_SECTION_PATH.exists():
        try:
            section_doc = json.loads(BI_PROFIT_SECTION_PATH.read_text(encoding="utf-8"))
            if prefer_section_payload(bi, section_doc, "profit"):
                bi["profit"] = extract_section(section_doc, "profit")
        except Exception:
            pass
    if BI_INVENTORY_SECTION_PATH.exists():
        try:
            section_doc = json.loads(BI_INVENTORY_SECTION_PATH.read_text(encoding="utf-8"))
            if prefer_section_payload(bi, section_doc, "inventoryDepletion"):
                bi["inventoryDepletion"] = extract_section(section_doc, "inventoryDepletion")
        except Exception:
            pass
    products = (bi.get("profit") or {}).get("products") or []
    profit_by_key: dict[str, dict] = {}
    for p in products:
        standard = clean(p.get("standard_goods_sn"))
        if not standard:
            continue
        for key in {standard, compact(standard)}:
            if key:
                profit_by_key[compact(key)] = p
    storage_by_key = build_storage_unit_map(bi)

    candidate_standards = set()
    candidate_standards.update(clean(p.get("standard_goods_sn")) for p in products if clean(p.get("standard_goods_sn")))
    candidate_standards.update(info["standard"] for info in storage_by_key.values() if clean(info.get("standard")))
    candidate_standards.update(cost_map.keys())

    true_map: dict[str, dict] = {}
    for standard in sorted(candidate_standards):
        standard = clean(standard)
        if not standard:
            continue
        keys = [standard, compact(standard)]
        p = None
        for key in keys:
            p = profit_by_key.get(compact(key))
            if p:
                break
        storage = None
        for key in keys:
            storage = storage_by_key.get(compact(key))
            if storage:
                break
        base_cost = lookup_cost(cost_map, keys)
        if base_cost is None and p:
            base_cost = num(p.get("unit_cost_sar"))
        if base_cost is None:
            continue
        storage_fee = num(storage.get("storageFeeSar")) if storage else (num(p.get("storage_fee_sar")) if p else None)
        storage_unit = num(storage.get("storageUnitCostSar")) if storage else None
        method = clean(storage.get("storageMethod")) if storage else (clean(p.get("storage_fee_method")) if p else "missing")
        true_unit = base_cost + storage_unit if storage_unit is not None else None
        info = {
            "unitCostSar": round(base_cost, 4),
            "storageUnitCostSar": None if storage_unit is None else round(storage_unit, 4),
            # Backward-compatible alias for old readers. Do not interpret this
            # as a 30-day value; the source of truth is storageUnitCostSar.
            "storageUnitCostSar30d": None if storage_unit is None else round(storage_unit, 4),
            "trueUnitCostSar": None if true_unit is None else round(true_unit, 4),
            "storageMethod": method,
            "storageFeeSar": None if storage_fee is None else round(storage_fee, 4),
            "quantityBasis": None if storage is None else storage.get("storageQuantityDays"),
            "storageCurrentQuantity": None if storage is None else storage.get("storageCurrentQuantity"),
            "storageUnitCostCandidateSar": None if storage is None else storage.get("storageUnitCostCandidateSar"),
            "storageBilledCurrentQuantity": None if storage is None else storage.get("storageBilledCurrentQuantity"),
            "storageOperationalSellableQuantity": None if storage is None else storage.get("storageOperationalSellableQuantity"),
            "storageOperationalDamagedQuantity": None if storage is None else storage.get("storageOperationalDamagedQuantity"),
            "storageOperationalPhysicalQuantity": None if storage is None else storage.get("storageOperationalPhysicalQuantity"),
            "storageOperationalInventorySnapshotDate": None if storage is None else storage.get("storageOperationalInventorySnapshotDate"),
            "storageOperationalStockPolicy": None if storage is None else storage.get("storageOperationalStockPolicy"),
            "storageAllocationQuantity": None if storage is None else storage.get("storageAllocationQuantity"),
            "storageAllocationQuantitySource": None if storage is None else storage.get("storageAllocationQuantitySource"),
            "storageQuantityEvidenceStatus": None if storage is None else storage.get("storageQuantityEvidenceStatus"),
            "storageQuantityDateGapDays": None if storage is None else storage.get("storageQuantityDateGapDays"),
            "storageQuantityRatioOperationalToBilled": None if storage is None else storage.get("storageQuantityRatioOperationalToBilled"),
            "storageQuantityRelativeDifference": None if storage is None else storage.get("storageQuantityRelativeDifference"),
            "storageNonpositiveQuantityFeeDays": None if storage is None else storage.get("storageNonpositiveQuantityFeeDays"),
            "storageInventoryCostBalanceSar": None if storage is None else storage.get("storageInventoryCostBalanceSar"),
            "storageAllHistoryUnitCostSar": None if storage is None else storage.get("storageAllHistoryUnitCostSar"),
            "storageRecent30UnitCostSar": None if storage is None else storage.get("storageRecent30UnitCostSar"),
            "storageRecent30FeeSar": None if storage is None else storage.get("storageRecent30FeeSar"),
            "storageRecent30Days": None if storage is None else storage.get("storageRecent30Days"),
            "storageAllDays": None if storage is None else storage.get("storageAllDays"),
            "storageUnitBasis": None if storage is None else storage.get("storageUnitBasis"),
            "storageSourceDateMin": None if storage is None else storage.get("storageSourceDateMin"),
            "storageSourceDateMax": None if storage is None else storage.get("storageSourceDateMax"),
            "source": "outputs/bi-portal/data.json",
        }
        for key in keys:
            if key:
                true_map[key] = info
    return true_map


def resolve_cost_input() -> Path:
    for candidate in COST_INPUT_CANDIDATES:
        if candidate.exists():
            return candidate
    expected = " / ".join(str(p) for p in COST_INPUT_CANDIDATES)
    raise SystemExit(f"missing cost workbook; expected one of: {expected}")


def main() -> None:
    from openpyxl import load_workbook

    input_path = resolve_cost_input()
    wb = load_workbook(input_path, data_only=True)
    ws = wb["成本"] if "成本" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise SystemExit("empty cost sheet")
    headers = [clean(v) for v in rows[0]]

    def idx(names: list[str]) -> int | None:
        keys = [n.replace("\n", "").replace(" ", "") for n in names]
        for i, h in enumerate(headers):
            hh = h.replace("\n", "").replace(" ", "")
            if hh in keys:
                return i
        for i, h in enumerate(headers):
            hh = h.replace("\n", "").replace(" ", "")
            if any(k in hh or hh in k for k in keys):
                return i
        return None

    model_i = idx(["型号", "货号", "标准货号"])
    name_i = idx(["希音标准名", "品名", "商品名称"])
    qty_i = idx(["数量", "发货数量", "总数量"])
    unit_i = idx(["单台总成本（SAR）", "单台总成本SAR", "单台总成本"])
    first_leg_i = idx(["头程运输费", "头程运输费金额"])

    batches = []
    agg: dict[str, dict] = {}
    for rno, row in enumerate(rows[1:], start=2):
        model = clean(row[model_i]) if model_i is not None else ""
        name = clean(row[name_i]) if name_i is not None else ""
        qty = num(row[qty_i]) if qty_i is not None else None
        unit_cost = num(row[unit_i]) if unit_i is not None else None
        first_leg = num(row[first_leg_i]) if first_leg_i is not None else None
        if not model or not qty or qty <= 0 or unit_cost is None:
            continue
        # 用户要求：没有头程运输费金额的批次先忽略，避免失真。
        if first_leg is None:
            continue
        full = f"{model}{name}".replace(" ", "")
        rec = agg.setdefault(full, {"quantity": 0.0, "total_cost": 0.0, "models": set(), "names": set()})
        rec["quantity"] += qty
        rec["total_cost"] += unit_cost * qty
        rec["models"].add(model)
        if name:
            rec["names"].add(name)
        batches.append({"row": rno, "model": model, "name": name, "full": full, "qty": qty, "unit_cost_sar": unit_cost})

    cost_map = {}
    for full, rec in agg.items():
        avg = rec["total_cost"] / rec["quantity"]
        keys = {full}
        for m in rec["models"]:
            keys.add(m)
        for m in rec["models"]:
            for n in rec["names"]:
                keys.add(f"{m}{n}".replace(" ", ""))
        for k in keys:
            cost_map[k] = round(avg, 4)

    true_cost_map = build_true_cost_map(cost_map)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps({
        "source": str(input_path),
        "biSource": str(BI_PATH) if BI_PATH.exists() else None,
        "biProfitSectionSource": str(BI_PROFIT_SECTION_PATH) if BI_PROFIT_SECTION_PATH.exists() else None,
        "count": len(cost_map),
        "trueCostCount": len(true_cost_map),
        "costMap": cost_map,
        "trueCostMap": true_cost_map,
        "batches": batches,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(OUTPUT), "count": len(cost_map), "trueCostCount": len(true_cost_map), "batches": len(batches)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
