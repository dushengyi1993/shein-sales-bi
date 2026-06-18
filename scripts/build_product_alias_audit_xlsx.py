from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
REPORTS_DIR = ROOT / "outputs" / "reports"
OUTPUT_PATH = REPORTS_DIR / "product-alias-audit-latest.xlsx"


def latest(prefix: str) -> Path:
    matches = sorted(REPORTS_DIR.glob(f"{prefix}*.csv"))
    if not matches:
        raise FileNotFoundError(f"missing report csv: {prefix}*.csv")
    return matches[-1]


def read_csv_rows(file: Path) -> list[list[str]]:
    with file.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.reader(f))


def autosize(ws, max_widths: dict[int, int] | None = None) -> None:
    max_widths = max_widths or {}
    for col_idx in range(1, ws.max_column + 1):
        width = 10
        for cell in ws.iter_cols(min_col=col_idx, max_col=col_idx, values_only=True):
            for value in cell:
                if value is None:
                    continue
                width = max(width, min(len(str(value)) * 1.4 + 3, max_widths.get(col_idx, 70)))
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def style_table(ws, max_widths: dict[int, int] | None = None) -> None:
    header_fill = PatternFill("solid", fgColor="201B16")
    header_font = Font(name="Microsoft YaHei UI", size=10, bold=True, color="FFF7E8")
    body_font = Font(name="Microsoft YaHei UI", size=10, color="201B16")
    thin = Side(style="thin", color="E4D9C8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = body_font
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    autosize(ws, max_widths)


def add_csv_sheet(wb: Workbook, title: str, file: Path, note: str) -> None:
    ws = wb.create_sheet(title)
    for row in read_csv_rows(file):
        ws.append(row)
    if ws.max_row == 0:
        ws.append(["无数据"])
    style_table(ws, {1: 32, 2: 18, 3: 90, 4: 48})
    note_row = ws.max_row + 2
    ws.cell(note_row, 1, note)
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=min(4, ws.max_column))
    note_cell = ws.cell(note_row, 1)
    note_cell.fill = PatternFill("solid", fgColor="FFF7E8")
    note_cell.font = Font(name="Microsoft YaHei UI", size=10, italic=True, color="6B6257")
    note_cell.alignment = Alignment(wrap_text=True, vertical="top")


def build() -> Path:
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    sources = [
        ("已确认别名", latest("product-alias-audit-"), "来自 config/product_aliases.json / product_catalog.json；这是当前已确认 canonical 规则。"),
        ("当前云端原始写法", latest("product-observed-aliases-"), "来自当前 BI linksData/data.json 可见原始写法；用于发现后台新写法是否需要归并。"),
        ("疑似可归并", latest("product-alias-candidates-"), "按型号和中文品类相似度生成，需人工确认后再写入别名配置。"),
        ("忽略项", latest("product-ignored-aliases-"), "认证标准号、噪音字段等应忽略，不作为标准货号。"),
    ]
    wb = Workbook()
    summary = wb.active
    summary.title = "说明"
    summary.sheet_view.showGridLines = False
    summary.append(["SHEIN BI V2 标准货号 / 别名审计表", "", "", ""])
    summary.merge_cells("A1:D1")
    summary["A1"].fill = PatternFill("solid", fgColor="201B16")
    summary["A1"].font = Font(name="Microsoft YaHei UI", size=16, bold=True, color="FFF7E8")
    summary["A1"].alignment = Alignment(vertical="center", wrap_text=True)
    summary.append([])
    rows = [
        ["工作表", "用途", "是否可直接归并", "说明"],
        ["已确认别名", "当前已写入配置的标准货号与别名", "是", "这里的规则会进入标准货号归并口径。"],
        ["当前云端原始写法", "线上可见的原始货号写法及出现次数", "否", "用于发现新别名，不能自动全部归并。"],
        ["疑似可归并", "按型号/中文相似度找出的候选", "需确认", "确认后再写入 config/product_aliases.json。"],
        ["忽略项", "认证标准号/噪音字段", "否", "例如 EN 62368 这类不应当作货号。"],
        ["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "", "基于当前 worktree 可用的 BI 输出快照。"],
    ]
    for row in rows:
        summary.append(row)
    style_table(summary, {1: 20, 2: 38, 3: 18, 4: 60})
    summary.freeze_panes = "A4"
    for title, file, note in sources:
        add_csv_sheet(wb, title, file, note)
    wb.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build())
